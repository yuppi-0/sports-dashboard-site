"""
WhoScoredの試合別イベントデータ（soccerdata経由）から、選手の「シーズン集計」と「試合別の表」を作る。

使い方:
    pip install soccerdata pandas numpy openpyxl socceraction   # socceraction はボール運び（プログレッシブラン）用。無くても動く
    python whoscored_player_stats.py --leagues Premier --seasons 2025 --n 3      # 動作確認（3試合だけ）
    python whoscored_player_stats.py --leagues Premier --seasons 2025 2026       # 全試合
    python whoscored_player_stats.py --leagues Premier LaLiga --seasons 2025 --out 出力

引数:
    --leagues  リーグ名（Premier / LaLiga / Bundesliga / SerieA / Ligue1。部分一致・複数可）。省略で全リーグ
    --seasons  シーズン（開幕年。2025 / 2025-26 のどちらでも可）。省略で最新シーズン
    --n        取得する試合数（省略で消化済みの全試合）。--last を付けると直近のn試合
    --rebuild  既存のファイルを使わず、最初から作り直す（取得済みの試合は保存分を使うので速い）
    --no-understat  Understat（xG・xA）を結合しない
    --no-approx     近似を含む指標（sca_*, gca, carries 系）を作らない
    --no-combine    取得のあとに、全リーグの結合をしない（並列実行用）
    --combine-only  取得はせず、結合だけを行う
    --out      出力フォルダ（省略時は、このスクリプトと同じ場所の output フォルダ）

続きから追加する（シーズン途中のリーグ向け）:
    出力ファイルがすでにあれば、その matches シートに「まだ入っていない終了済みの試合」だけを足す。
    取得済みの試合は取り直さない。players シートは、足したあとの matches から作り直す。
    集計ロジックを変えた版のファイルは、自動で作り直される（--rebuild で手動でも可）。
    試合中・未開催の試合は取らない（終了した試合だけ）。

出力（--out の下。省略時はスクリプトと同じ場所の output フォルダ）:
    <リーグ名>/<リーグ名>_<開幕年>.xlsx        例: LaLiga/LaLiga_2025.xlsx
        players シート: 選手のシーズン集計（venue 列が all / home / away の3種類。90分あたり列つき）
        matches シート: 選手×試合の表（日付・対戦相手・ホーム/アウェイ・出場時間・評価・各指標）
    結合/<開幕年>年/all_leagues_players_<開幕年>.xlsx   全リーグの players を、シーズンごとに結合（1シート）
    結合/<開幕年>年/all_leagues_matches_<開幕年>.xlsx   全リーグの matches を、シーズンごとに結合（同上）。実行のたびに作り直す
    python whoscored_player_stats.py --combine-only   … 取得せず結合だけやり直す
    ※ 並列で実行するときは、各プロセスに --no-combine を付け、全部終わってから --combine-only を1回だけ実行する

指標について:
    - 試合ごとのイベントから数える指標: シュート（枠内・ブロック・ペナルティエリア内・ヘディング）、タッチ（ゾーン別・ペナルティエリア内）、
      パス（距離別・ライブ/デッドボール・FK・スルーパス・サイドチェンジ・クロス・アシスト）、守備（ゾーン別のタックル・インターセプト等）、
      GK（セーブ・失点・クリーンシート・キャッチ・パンチ・スイーパー）。
    - 近似を含む指標（sca_*, gca, carries 系）: シュートにつながったプレー（SCA/GCA）は、シュート直前の同じチームのプレーを
      最大2つまでさかのぼって数える。ボール運びは、連続する2つのプレーの間を補完したもので、実際の運びとはずれる。
      使わない場合は RANK_SPECS から該当の指標を削除する（列自体は、データ加工で落とせる）。
    - 順位（<指標>_順位・<指標>_順位_母数）は、RANK_SPECS の指標について、同じリーグ・シーズン・集団の中で付ける。

注意:
- 座標はWhoScored形式（0-100。x は自陣ゴール→相手ゴール方向に増加）。
- プログレッシブパス等の定義は自作のため、FBrefの値とは一致しない。
- 評価（rating）・先発・出場時間は、soccerdata が手元に保存している試合ごとのJSON
  （~/soccerdata/data/WhoScored/events/…）から読む。JSONの構造が想定と違うと空欄になる。
- 初回は表示される qualifier 名（KeyPass, Cross 等）が実データと一致しているか確認すること。
"""
import argparse
import json
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd

DEFAULT_OUT = Path(__file__).resolve().parent / "output"   # 既定の出力先: このスクリプトと同じ場所の output フォルダ

# ---- 定義（必要に応じて調整） -------------------------------------------
FINAL_THIRD_X = 66.7                                 # ファイナルサード開始位置
BOX_X, BOX_Y_LO, BOX_Y_HI = 84.3, 20.4, 79.6        # ペナルティエリア（相手側）
PROGRESSIVE_MIN_DX = 10.0                            # 前進距離の下限（座標単位≒約10m）
DEF_THIRD_X = 33.3                                   # 自陣3分の1の終わり
PITCH_L, PITCH_W = 105.0, 68.0                       # ピッチの長さ・幅（m）。WhoScoredの座標(0-100)をmに直すのに使う
SHORT_PASS_M = (4.57, 13.72)                         # 短いパス: 5〜15ヤード
MEDIUM_PASS_M = (13.72, 27.43)                       # 中距離: 15〜30ヤード（これ以上は長距離）
SWITCH_MIN_LATERAL_M = 36.6                          # サイドチェンジ: 横方向に40ヤード以上
SCA_MAX_GAP_S = 15                                   # シュートにつながるプレーとして数える、最大の時間差（秒）
SCA_KINDS = ["pass_live", "pass_dead", "take_on", "shot", "foul_won", "defensive"]
SCA_BREAK_OPP = {"Pass", "Interception", "BallRecovery", "Clearance", "TakeOn", "KeeperPickup", "Claim"}   # 出たら連なりを止める相手のプレー
CARRY_PROGRESSIVE_MIN_M = 10.0                       # プログレッシブラン: 前進10m以上（自陣40%で終わるものは除く）
MAX_MATCH_MINUTES = 90                               # 出場時間の上限（リーグ戦）
MIN_MINUTES_P90 = 10                                 # これ未満の出場時間では90分あたりを出さない

QUALIFIERS = ["Longball", "Cross", "Throughball", "KeyPass",
              "ThrowIn", "GoalKick", "CornerTaken", "FreekickTaken", "OwnGoal",
              "Head", "Penalty"]
OWN_GOAL_QUALIFIER_ID = 28                           # Optaのqualifier番号（オウンゴール）

CODE_VERSION = 6                                     # 集計ロジックを変えたら上げる（既存ファイルは自動で作り直される）
VERSION_NOTE = f"whoscored_player_stats v{CODE_VERSION}"

COUNT_COLS = [  # aggregate() がイベントから数える列（この並びで出力する）
    "passes", "passes_completed", "progressive_passes", "passes_final_third", "passes_into_box",
    "key_passes", "crosses", "crosses_completed", "long_balls", "long_balls_completed",
    "through_balls", "through_balls_completed",
    "passes_short", "passes_short_completed", "passes_medium", "passes_medium_completed",
    "passes_long", "passes_long_completed",
    "passes_dead", "passes_dead_completed", "passes_fk", "passes_fk_completed",
    "switches", "switches_completed",
    "take_ons", "take_ons_won", "dispossessed", "errors",
    "tackles", "tackles_won", "interceptions", "clearances", "blocked_passes", "ball_recoveries",
    "tackles_def_third", "tackles_mid_third", "tackles_att_third",
    "tackles_won_def_third", "tackles_won_mid_third", "tackles_won_att_third",
    "interceptions_def_third", "interceptions_mid_third", "interceptions_att_third",
    "clearances_def_third", "clearances_mid_third", "clearances_att_third",
    "ball_recoveries_def_third", "ball_recoveries_mid_third", "ball_recoveries_att_third",
    "def_actions_def_third", "def_actions_mid_third", "def_actions_att_third",
    "aerials", "aerials_won", "fouls_committed", "fouls_won", "dribbled_past",
    "shots", "shots_on_target", "shots_off_target", "shots_blocked", "shots_in_box", "headed_shots",
    "goals", "own_goals",
    "touches", "touches_def_third", "touches_mid_third", "touches_att_third", "touches_att_pen",
    "claims", "punches", "keeper_pickups", "gk_sweeper_actions",
    "save_events",          # GKのセーブと、フィールドプレーヤーのシュートブロックが同じ「Save」で記録されるため、後で分ける
]
# 試合別の表を作る段階で足す列（SCA・ボール運び・GK）
EXTRA_COUNT_COLS = (
    [f"sca_{k}" for k in SCA_KINDS] + ["sca", "gca"]
    + ["carries", "carry_distance", "progressive_carries", "carries_final_third", "carries_into_box"]
    + ["saves", "shot_blocks", "goals_conceded", "clean_sheets"]
)

POS_GROUP = {   # WhoScoredの試合ごとのポジション -> FW/MF/DF/GK（交代で入った選手の "Sub" は対象外）
    "GK": "GK", "DC": "DF", "DL": "DF", "DR": "DF",
    "DMC": "MF", "DML": "MF", "DMR": "MF", "MC": "MF", "ML": "MF", "MR": "MF",
    "AMC": "MF", "AML": "MF", "AMR": "MF", "FW": "FW", "FWL": "FW", "FWR": "FW",
}
SIDE_BACK_POSITIONS = {"DL", "DR"}          # DFのうち、これらでの出場時間が半分以上ならSB、そうでなければCB
SB_MIN_SHARE = 0.5

RANK_MIN_MINUTES = 900                       # 順位を出す条件: シーズン通算の出場時間
# 順位を付ける指標: 集団 -> [(列名, high=大きいほど良い / low=小さいほど良い, 分母の列, 分母の下限)]
# 集団は FW / MF / CB / SB / GK（DFはCBとSBに分ける）。順位は「同じリーグ・同じシーズン・同じ集団」の中で付ける。
# 分母の列は、その指標の試行数（例: タックル成功率ならタックル数）。通算の値が下限未満の選手は順位を付けない。
# ★印は近似を含む指標（SCA=シュートにつながったプレー、ボール運び）。使わない場合はここから削除する。
_ZONES = ("def_third", "mid_third", "att_third")
_DF_ZONE_SPECS = (
    [(f"tackles_{z}_p90", "high", None, 0) for z in _ZONES]
    + [(f"tackle_win_pct_{z}", "high", f"tackles_{z}", 5) for z in _ZONES]
    + [(f"interceptions_{z}_p90", "high", None, 0) for z in _ZONES]
)
RANK_SPECS = {
    "FW": [  # 付録A: コア指標
           ("goals_p90", "high", None, 0), ("xg_diff_p90", "high", None, 0),
           ("shot_on_target_pct", "high", "shots", 20), ("touches_att_pen_pct", "high", "touches", 100),
           ("take_on_pct", "high", "take_ons", 10), ("def_actions_att_third_p90", "high", None, 0),
           # 付録B: チャンス創出の起点別 ★
           ("sca_pass_live_p90", "high", "sca_pass_live", 5), ("sca_pass_dead_p90", "high", "sca_pass_dead", 3),
           ("sca_take_on_p90", "high", "sca_take_on", 3), ("sca_shot_p90", "high", "sca_shot", 3),
           ("sca_foul_won_p90", "high", "sca_foul_won", 3), ("sca_defensive_p90", "high", "sca_defensive", 3)],
    "MF": [  # 付録A
           ("pass_pct", "high", "passes", 100), ("progressive_passes_p90", "high", None, 0),
           ("us_xa_p90", "high", None, 0), ("key_passes_p90", "high", None, 0),
           ("tackle_win_pct", "high", "tackles", 10), ("interceptions_p90", "high", None, 0),
           # 付録B: パス種類別
           ("pass_short_pct", "high", "passes_short", 50), ("pass_medium_pct", "high", "passes_medium", 50),
           ("pass_long_pct", "high", "passes_long", 20), ("passes_dead_pct", "high", "passes_dead", 10),
           ("passes_fk_pct", "high", "passes_fk", 5), ("through_ball_pct", "high", "through_balls", 5),
           ("switch_pct", "high", "switches", 10), ("cross_pct", "high", "crosses", 10)],
    "CB": [  # 付録A
           ("tackle_win_pct", "high", "tackles", 10), ("interceptions_p90", "high", None, 0),
           ("clearances_p90", "high", None, 0), ("aerial_win_pct", "high", "aerials", 10),
           ("dribbled_past_pct", "low", "challenges", 10), ("pass_pct", "high", "passes", 100),
           ] + _DF_ZONE_SPECS,                                   # 付録B: ゾーン別守備
    "SB": [  # 付録A
           ("tackle_win_pct", "high", "tackles", 10), ("dribbled_past_pct", "low", "challenges", 10),
           ("crosses_p90", "high", None, 0), ("progressive_carries_p90", "high", None, 0),   # ★
           ("key_passes_p90", "high", None, 0), ("duel_win_pct", "high", "duels", 20),
           ] + _DF_ZONE_SPECS,
    "GK": [  # 付録A
           ("gk_save_pct", "high", "shots_faced", 20), ("goals_conceded_p90", "low", None, 0),
           ("clean_sheet_pct", "high", "matches", 10), ("pass_pct", "high", "passes", 100),
           ("long_ball_pct", "high", "long_balls", 10), ("gk_sweeper_actions_p90", "high", None, 0)],
}

SD_LEAGUES = {  # 短いリーグ名 -> soccerdata のリーグID
    "Premier": "ENG-Premier League",
    "LaLiga": "ESP-La Liga",
    "Bundesliga": "GER-Bundesliga",
    "SerieA": "ITA-Serie A",
    "Ligue1": "FRA-Ligue 1",
}


# ---- リーグ・シーズンの指定 ----------------------------------------------------
def latest_season_year(today=None):
    t = today or pd.Timestamp.now()
    return t.year if t.month >= 7 else t.year - 1


def season_label(year):
    """2025 -> '2025-2026'"""
    return f"{year}-{year + 1}"


def sd_season(year):
    """2025 -> '2025-26'（soccerdata の表記）"""
    return f"{year}-{str(year + 1)[-2:]}"


def pick_leagues(specs):
    if not specs:
        return list(SD_LEAGUES)
    out = []
    for sp in specs:
        hits = [k for k in SD_LEAGUES if sp.lower() in f"{k} {SD_LEAGUES[k]}".lower()]
        if not hits:
            raise SystemExit(f"リーグ '{sp}' が見つかりません。候補: {list(SD_LEAGUES)}")
        out += [h for h in hits if h not in out]
    return out


def parse_years(specs):
    if not specs:
        return [latest_season_year()]
    years = []
    for s in specs:
        m = re.search(r"(\d{4})", str(s))
        if not m:
            raise SystemExit(f"シーズン '{s}' を解釈できません（例: 2025 / 2025-26）")
        years.append(int(m[1]))
    return years


def build_jobs(league_specs, season_specs):
    """(リーグ短縮名, 開幕年) の一覧。"""
    return [(name, y) for name in pick_leagues(league_specs) for y in parse_years(season_specs)]


# ---- 取得 ------------------------------------------------------------------
def select_match_ids(schedule, n=None, last=False):
    """
    日程から「終了した試合」の game_id を選ぶ（試合中・未開催の試合は含めない）。
    n=None : 終了した全試合 / n=5 : 開幕から5試合（last=True なら直近5試合）
    """
    s = _flat(schedule)
    now = pd.Timestamp.now(tz="UTC")
    s = s.assign(_date=pd.to_datetime(s["date"], utc=True, errors="coerce"))
    started = s[s["_date"] < now]
    done = pd.DataFrame()
    if "elapsed" in s.columns:                       # 終了した試合は elapsed が FT（延長・PKは AET・PEN）
        done = started[started["elapsed"].astype(str).str.upper().isin(["FT", "AET", "PEN"])]
    if done.empty:                                   # 列が無い／該当なしなら、開始から3時間以上たった試合
        done = s[s["_date"] < now - pd.Timedelta(hours=3)]
        if "elapsed" in s.columns and not done.empty:
            print("  注意: 試合の終了状態を判定できなかったため、開始から3時間以上たった試合を対象にします")
    ids = done.sort_values("_date")["game_id"].astype(int).tolist()
    if n is None:
        return ids
    return ids[-n:] if last else ids[:n]


def missing_games(events, ids):
    """取得を試みた試合IDのうち、イベントが1件も返ってこなかったものを返す。"""
    try:
        got = set(pd.to_numeric(_flat(events)["game_id"], errors="coerce").dropna().astype(int))
    except KeyError:
        got = set()
    return sorted(set(int(i) for i in ids) - got)


def fetch_carries(ws, ids):
    """
    保存済みの試合JSONから、ボール運び（SPADLのdribble）を作って回数を数える。
    socceraction が入っていなければ None を返す（ボール運びの列は空欄になる）。
    """
    try:
        spadl = ws.read_events(match_id=ids, force_cache=True, output_fmt="spadl", on_error="skip")
        return compute_carries(spadl)
    except ImportError:
        print("  注意: socceraction が無いため、ボール運びの指標をスキップします"
              "（python -m pip install socceraction で入ります）")
    except Exception as e:  # noqa: BLE001
        print(f"  注意: ボール運びの計算をスキップします: {e}")
    return None


def fetch_events(league, year, n=None, last=False, skip_ids=(), with_carries=True):
    """
    WhoScoredから試合のイベントと日程を取得する。skip_ids（取得済みの試合ID）は取らない。
    戻り値: (events, schedule, match_files, carries)。新しく取る試合が無ければ events は None。
    match_files は {game_id: 手元に保存された試合JSONのパス}、carries は選手×試合のボール運び（無ければ None）。
    """
    import soccerdata as sd  # 集計だけ試す場合にインポート不要にするため遅延

    ws = sd.WhoScored(leagues=league, seasons=sd_season(year))
    schedule = ws.read_schedule()
    all_ids = select_match_ids(schedule, n=n, last=last)
    if not all_ids:
        raise RuntimeError("取得対象の試合が見つかりません（日程に終了した試合がない）")
    skip = {int(i) for i in skip_ids}
    ids = [i for i in all_ids if i not in skip]
    print(f"終了した{len(all_ids)}試合のうち、新しく取得するのは{len(ids)}試合です（既存{len(all_ids) - len(ids)}試合）")
    if not ids:
        return None, schedule, {}, None
    events = ws.read_events(match_id=ids, force_cache=True, on_error="skip")  # 日程は取得済みなので再取得しない
    missing = missing_games(events, ids)
    if missing:
        print(f"  警告: {len(missing)}試合を取得できませんでした（ブロックや通信エラーの可能性）。"
              "同じコマンドを再実行すると、不足分だけ取得します")
    sch = _flat(schedule)
    wanted = set(ids)
    files = {
        int(r.game_id): Path(ws.data_dir) / "events" / f"{r.league}_{r.season}" / f"{int(r.game_id)}.json"
        for r in sch.itertuples() if int(r.game_id) in wanted
    }
    carries = fetch_carries(ws, ids) if with_carries else None
    return events, schedule, files, carries


def fetch_understat(league, year):
    """Understatの選手シーズン成績（xG, xA, 出場時間など）を取得する。"""
    import soccerdata as sd

    return sd.Understat(leagues=league, seasons=sd_season(year)).read_player_season_stats()


# ---- 試合ごとのJSON（評価・先発・出場時間） -----------------------------------------
def _last_rating(r):
    try:
        if isinstance(r, dict) and r:
            return float(r[max(r, key=lambda x: float(x))])
        if isinstance(r, (list, tuple)) and r:
            return float(r[-1])
    except (TypeError, ValueError):
        pass
    return np.nan


def _red_card_minutes(team):
    """退場（レッド・2枚目のイエロー）した選手の分（延長込み）。カードは team の incidentEvents にある。"""
    out = {}
    for e in team.get("incidentEvents") or []:
        if (isinstance(e, dict) and e.get("playerId") is not None
                and _name(e.get("cardType")) in ("Red", "SecondYellow")
                and e.get("expandedMinute") is not None):
            out.setdefault(e["playerId"], e["expandedMinute"])
    return out


def _minutes(p, has_lineup, total, red):
    """
    出場時間を推定する。soccerdata/socceraction と同じく、延長込みの分（expanded minute）で
    出場した割合を出し、90分換算にする（試合の長さに左右されない）。
    has_lineup: その試合のJSONに先発情報（isFirstEleven）があるか。ベンチ発の選手はこのキーが無い。
    途中出場なのに入った分が無い場合は不明（NaN）にする。
    """
    if not has_lineup:
        return np.nan
    t_out = red.get(p["playerId"], p.get("subbedOutExpandedMinute"))
    end = total if t_out is None else min(t_out, total)
    if p.get("isFirstEleven"):
        start = 0
    else:
        start = p.get("subbedInExpandedMinute")
        if start is None:
            return np.nan
    return int(round(max(0, end - start) * MAX_MATCH_MINUTES / total))


def read_match_players(match_files):
    """試合JSONから、選手ごとの先発・出場時間・評価・試合最優秀選手を読む。読めなければ空のDataFrame。"""
    rows = []
    for gid, path in match_files.items():
        try:
            d = json.loads(Path(path).read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if not isinstance(d, dict):
            continue
        teams = [d.get(side) or {} for side in ("home", "away")]
        total = max(d.get("expandedMaxMinute") or d.get("maxMinute") or 0, MAX_MATCH_MINUTES)
        for team in teams:
            players = [p for p in team.get("players") or []
                       if isinstance(p, dict) and p.get("playerId") is not None]
            has_lineup = any("isFirstEleven" in p for p in players)
            red = _red_card_minutes(team)
            for p in players:
                rows.append({
                    "game_id": int(gid),
                    "player_id": int(p["playerId"]),
                    "position_played": p.get("position"),
                    "is_starter": bool(p.get("isFirstEleven")) if has_lineup else None,
                    "man_of_match": bool(p.get("isManOfTheMatch")) if has_lineup else None,
                    "minutes": _minutes(p, has_lineup, total, red),
                    "rating": _last_rating((p.get("stats") or {}).get("ratings")),
                })
    return pd.DataFrame(rows)


# ---- 前処理 ----------------------------------------------------------------
def _name(v):
    return v.get("displayName") if isinstance(v, dict) else v


def _qualifier_names(quals):
    if not isinstance(quals, (list, tuple)):
        return frozenset()
    names = set()
    for q in quals:
        if isinstance(q, dict):
            t = q.get("type")
            n = _name(t)
            if n:
                names.add(n)
            if isinstance(t, dict) and t.get("value") == OWN_GOAL_QUALIFIER_ID:
                names.add("OwnGoal")                  # 表示名が違っても番号で拾う
    return frozenset(names)


def check_qualifiers(events, top=40):
    """データ中に実在するイベント種類名とqualifier名の上位を表示する（名前の確認用）。"""
    print(events["type"].map(_name).value_counts().head(60).to_string())
    names = events["qualifiers"].map(_qualifier_names).explode()
    print(names.value_counts().head(150).to_string())


def _flat(df):
    if isinstance(df.index, pd.MultiIndex) or df.index.name is not None:
        df = df.reset_index()
    return df.copy()


def _in_box(x, y):
    return (x >= BOX_X) & (y >= BOX_Y_LO) & (y <= BOX_Y_HI)


def prepare(events):
    d = _flat(events)
    d["type"] = d["type"].map(_name)
    d["ok"] = d["outcome_type"].map(_name).eq("Successful")
    qs = d["qualifiers"].map(_qualifier_names)
    for name in QUALIFIERS:
        d["q_" + name] = qs.map(lambda s, n=name: n in s)
    for c in ["x", "y", "end_x", "end_y"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    for c in ["is_shot", "is_goal", "is_touch"]:
        d[c] = d[c].fillna(False).astype(bool) if c in d else False
    return d


def add_match_info(events, schedule):
    """
    イベントに date / venue（home・away）/ opponent（と得点）を付ける。日程と突き合わせる。
    ホーム/アウェイはチームIDで判定する（日程とイベントでチーム名の表記が違うことがあるため）。
    """
    ev = _flat(events)
    sch = _flat(schedule)
    id_cols = [c for c in ("home_team_id", "away_team_id") if c in sch.columns]
    score_cols = [c for c in ("home_score", "away_score") if c in sch.columns]
    keep = ["league", "season", "game", "date", "home_team", "away_team"] + id_cols + score_cols
    d = ev.merge(sch[keep], on=["league", "season", "game"], how="left")

    if len(id_cols) == 2 and "team_id" in d.columns:
        tid = pd.to_numeric(d["team_id"], errors="coerce")
        hid = pd.to_numeric(d["home_team_id"], errors="coerce")
        aid = pd.to_numeric(d["away_team_id"], errors="coerce")
        is_home, is_away = tid == hid, tid == aid
    else:                                            # チームIDが無ければ名前で判定（表記が違うと外れる）
        is_home, is_away = d["team"] == d["home_team"], d["team"] == d["away_team"]
    d["venue"] = np.select([is_home, is_away], ["home", "away"], default="")
    d["venue"] = d["venue"].replace("", np.nan)

    # 対戦相手の名前は、イベント側のチーム名に揃える（日程側の表記とずれるため）
    if len(id_cols) == 2 and "team_id" in d.columns:
        names = {(g, t): n for g, t, n in
                 d[["game_id", "team_id", "team"]].dropna().drop_duplicates().itertuples(index=False)}
        opp_id = np.where(is_home, aid, np.where(is_away, hid, np.nan))
        d["opponent"] = [names.get((g, o)) if pd.notna(o) else None
                         for g, o in zip(d["game_id"], opp_id)]
    else:
        d["opponent"] = np.where(is_home, d["away_team"], np.where(is_away, d["home_team"], None))

    d["date"] = pd.to_datetime(d["date"], utc=True, errors="coerce").dt.strftime("%Y-%m-%d")
    if len(score_cols) == 2:
        d["goals_for"] = np.where(is_home, d["home_score"], np.where(is_away, d["away_score"], np.nan))
        d["goals_against"] = np.where(is_home, d["away_score"], np.where(is_home | is_away, d["home_score"], np.nan))
    n_bad = int(d["venue"].isna().sum())
    if n_bad:
        print(f"  注意: {n_bad}行でホーム/アウェイを判定できませんでした（venue・opponent・得点が空欄になります）")
    return d.drop(columns=["home_team", "away_team"] + id_cols + score_cols)


# ---- 集計 ------------------------------------------------------------------
def _add_rates(out):
    """回数の列から成功率などを計算する（out は集計済みの回数の列を含む DataFrame。無い列に依存するものは飛ばす）。"""
    new = {}

    def pct(a, b):
        return (out[a] / out[b].replace(0, np.nan) * 100).round(1)

    def has(*cols):
        return all(c in out.columns for c in cols)

    new["pass_pct"] = pct("passes_completed", "passes")
    new["cross_pct"] = pct("crosses_completed", "crosses")
    new["long_ball_pct"] = pct("long_balls_completed", "long_balls")
    new["take_on_pct"] = pct("take_ons_won", "take_ons")
    new["tackle_win_pct"] = pct("tackles_won", "tackles")
    new["aerial_win_pct"] = pct("aerials_won", "aerials")
    challenges = out["tackles"] + out["dribbled_past"]                 # 1対1の場面（タックル＋かわされた）
    new["dribbled_past_pct"] = (out["dribbled_past"] / challenges.replace(0, np.nan) * 100).round(1)
    duels = out["tackles"] + out["aerials"]                            # デュエル（タックル＋空中戦）
    new["duel_win_pct"] = ((out["tackles_won"] + out["aerials_won"]) / duels.replace(0, np.nan) * 100).round(1)
    if has("shots_on_target", "shots"):
        new["shot_on_target_pct"] = pct("shots_on_target", "shots")
    if has("touches_att_pen", "touches"):
        new["touches_att_pen_pct"] = pct("touches_att_pen", "touches")
    for k in ("short", "medium", "long"):
        if has(f"passes_{k}_completed", f"passes_{k}"):
            new[f"pass_{k}_pct"] = pct(f"passes_{k}_completed", f"passes_{k}")
    if has("passes_dead_completed", "passes_dead"):
        new["passes_dead_pct"] = pct("passes_dead_completed", "passes_dead")
    if has("passes_fk_completed", "passes_fk"):
        new["passes_fk_pct"] = pct("passes_fk_completed", "passes_fk")
    if has("switches_completed", "switches"):
        new["switch_pct"] = pct("switches_completed", "switches")
    if has("through_balls_completed", "through_balls"):
        new["through_ball_pct"] = pct("through_balls_completed", "through_balls")
    for z in ("def_third", "mid_third", "att_third"):
        if has(f"tackles_won_{z}", f"tackles_{z}"):
            new[f"tackle_win_pct_{z}"] = pct(f"tackles_won_{z}", f"tackles_{z}")
    if has("saves", "goals_conceded"):                                 # GK: 枠内シュートに対するセーブの割合
        faced = out["saves"] + out["goals_conceded"]
        new["gk_save_pct"] = (out["saves"] / faced.replace(0, np.nan) * 100).round(1)
    if has("clean_sheets", "matches"):
        new["clean_sheet_pct"] = pct("clean_sheets", "matches")
    return pd.concat([out, pd.DataFrame(new, index=out.index)], axis=1)


def aggregate(events, by=("league", "season", "team", "player")):
    """選手別（by の粒度）にスタッツを集計する。by に game_id 等を入れれば試合別になる。"""
    d = prepare(events)
    d = d[d["player"].notna()].copy()
    d["_g"] = d["game"]
    by = list(by)
    t, ok = d["type"], d["ok"]

    is_pass = t.eq("Pass")
    set_piece = d["q_ThrowIn"] | d["q_GoalKick"] | d["q_CornerTaken"] | d["q_FreekickTaken"]
    open_pass = is_pass & ~set_piece
    dx = d["end_x"] - d["x"]
    start_in_box = _in_box(d["x"], d["y"])
    end_in_box = _in_box(d["end_x"], d["end_y"])
    in_box = _in_box(d["x"], d["y"])
    def_action = t.isin(["Tackle", "Interception", "BallRecovery"])

    # パスの距離（m）
    dist = np.hypot((d["end_x"] - d["x"]) * PITCH_L / 100, (d["end_y"] - d["y"]) * PITCH_W / 100)
    short = is_pass & (dist >= SHORT_PASS_M[0]) & (dist < SHORT_PASS_M[1])
    medium = is_pass & (dist >= MEDIUM_PASS_M[0]) & (dist < MEDIUM_PASS_M[1])
    long_ = is_pass & (dist >= MEDIUM_PASS_M[1])
    lateral = (d["end_y"] - d["y"]).abs() * PITCH_W / 100
    switch = open_pass & (lateral >= SWITCH_MIN_LATERAL_M)
    dead = is_pass & set_piece
    fk = is_pass & d["q_FreekickTaken"]

    # ゾーン（自陣・中央・敵陣の3分の1）
    z_def = d["x"] < DEF_THIRD_X
    z_att = d["x"] >= FINAL_THIRD_X
    z_mid = d["x"].notna() & ~z_def & ~z_att

    # シュート
    is_own = d["q_OwnGoal"]
    blocked = t.eq("SavedShot") & (d["blocked_by_outfield"] if "blocked_by_outfield" in d else False)
    on_target = (t.eq("Goal") & ~is_own) | (t.eq("SavedShot") & ~blocked)
    off_target = t.isin(["MissedShots", "ShotOnPost"])
    shot = d["is_shot"] & ~is_own

    flags = pd.DataFrame({
        # パス
        "passes": open_pass,
        "passes_completed": open_pass & ok,
        "progressive_passes": open_pass & ok & (dx >= PROGRESSIVE_MIN_DX),
        "passes_final_third": open_pass & ok & (d["x"] < FINAL_THIRD_X) & (d["end_x"] >= FINAL_THIRD_X),
        "passes_into_box": open_pass & ok & ~start_in_box & end_in_box,
        "key_passes": is_pass & d["q_KeyPass"],
        "crosses": is_pass & d["q_Cross"] & ~d["q_CornerTaken"],
        "crosses_completed": is_pass & d["q_Cross"] & ~d["q_CornerTaken"] & ok,
        "long_balls": is_pass & d["q_Longball"],
        "long_balls_completed": is_pass & d["q_Longball"] & ok,
        "through_balls": is_pass & d["q_Throughball"],
        "through_balls_completed": is_pass & d["q_Throughball"] & ok,
        "passes_short": short, "passes_short_completed": short & ok,
        "passes_medium": medium, "passes_medium_completed": medium & ok,
        "passes_long": long_, "passes_long_completed": long_ & ok,
        "passes_dead": dead, "passes_dead_completed": dead & ok,
        "passes_fk": fk, "passes_fk_completed": fk & ok,
        "switches": switch, "switches_completed": switch & ok,
        # ドリブル・ボールロスト
        "take_ons": t.eq("TakeOn"),
        "take_ons_won": t.eq("TakeOn") & ok,
        "dispossessed": t.eq("Dispossessed"),
        "errors": t.eq("Error"),
        # 守備
        "tackles": t.eq("Tackle"),
        "tackles_won": t.eq("Tackle") & ok,
        "interceptions": t.eq("Interception"),
        "clearances": t.eq("Clearance"),
        "blocked_passes": t.eq("BlockedPass"),
        "ball_recoveries": t.eq("BallRecovery"),
        "tackles_def_third": t.eq("Tackle") & z_def,
        "tackles_mid_third": t.eq("Tackle") & z_mid,
        "tackles_att_third": t.eq("Tackle") & z_att,
        "tackles_won_def_third": t.eq("Tackle") & ok & z_def,
        "tackles_won_mid_third": t.eq("Tackle") & ok & z_mid,
        "tackles_won_att_third": t.eq("Tackle") & ok & z_att,
        "interceptions_def_third": t.eq("Interception") & z_def,
        "interceptions_mid_third": t.eq("Interception") & z_mid,
        "interceptions_att_third": t.eq("Interception") & z_att,
        "clearances_def_third": t.eq("Clearance") & z_def,
        "clearances_mid_third": t.eq("Clearance") & z_mid,
        "clearances_att_third": t.eq("Clearance") & z_att,
        "ball_recoveries_def_third": t.eq("BallRecovery") & z_def,
        "ball_recoveries_mid_third": t.eq("BallRecovery") & z_mid,
        "ball_recoveries_att_third": t.eq("BallRecovery") & z_att,
        "def_actions_def_third": def_action & z_def,
        "def_actions_mid_third": def_action & z_mid,
        "def_actions_att_third": def_action & z_att,
        "aerials": t.eq("Aerial"),
        "aerials_won": t.eq("Aerial") & ok,
        "fouls_committed": t.eq("Foul") & ~ok,
        "fouls_won": t.eq("Foul") & ok,
        "dribbled_past": t.eq("Challenge") & ~ok,
        # シュート
        "shots": shot,
        "shots_on_target": on_target & shot,
        "shots_off_target": off_target & shot,
        "shots_blocked": blocked & shot,
        "shots_in_box": shot & in_box,
        "headed_shots": shot & d["q_Head"],
        "goals": d["is_goal"] & ~is_own,               # オウンゴールは含めない
        "own_goals": d["is_goal"] & is_own,            # 自チームのゴールに入れてしまった数
        # タッチ
        "touches": d["is_touch"],
        "touches_def_third": d["is_touch"] & z_def,
        "touches_mid_third": d["is_touch"] & z_mid,
        "touches_att_third": d["is_touch"] & z_att,
        "touches_att_pen": d["is_touch"] & in_box,
        # GK専用のイベント
        "claims": t.eq("Claim"),
        "punches": t.eq("Punch"),
        "keeper_pickups": t.eq("KeeperPickup"),
        "gk_sweeper_actions": t.eq("KeeperSweeper"),
        "save_events": t.eq("Save"),
    })

    flags = flags[COUNT_COLS]                            # 列の並びを COUNT_COLS に固定する
    keys = [d[c] for c in by]
    out = flags.groupby(keys, dropna=False).sum().astype(int)
    out.insert(0, "matches", d["_g"].groupby(keys, dropna=False).nunique().to_numpy())

    return _add_rates(out).reset_index()


# ---- 試合別の表・選手のシーズン表 ---------------------------------------------------
def _int_ids(df):
    for c in ("game_id", "player_id"):
        if c in df:
            df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    return df


GK_ONLY_EVENT_TYPES = ("Claim", "Punch", "KeeperPickup", "KeeperSweeper")
STOP_PAIR_WINDOW = 4          # SavedShot と対応する Save のイベントを探す範囲（前後の件数）
STOP_PAIR_MAX_GAP_S = 3       # 同上（秒）


def _time_sorted(events):
    """イベントを、試合→時刻→イベント番号の順に並べる（prepare 済みの DataFrame を返す）。"""
    d = prepare(events)
    for c in ("expanded_minute", "minute", "second", "event_id"):
        if c not in d:
            d[c] = np.nan
    minute = pd.to_numeric(d["expanded_minute"], errors="coerce").fillna(pd.to_numeric(d["minute"], errors="coerce"))
    d["_t"] = minute * 60 + pd.to_numeric(d["second"], errors="coerce").fillna(0)
    d["_eid"] = pd.to_numeric(d["event_id"], errors="coerce")
    d = d.sort_values(["game_id", "_t", "_eid"], kind="mergesort")
    d["_period"] = (d["period"] if "period" in d else pd.Series("", index=d.index)).map(_name).astype(str)
    return d


def mark_shot_stops(events, match_players=None):
    """
    「SavedShot」（止められたシュート）を、GKが止めたもの（枠内）と、フィールドの選手がブロックしたものに分ける。
    WhoScoredでは、どちらも相手チームの「Save」として記録される。SavedShot の前後にある、
    相手チームの Save を1対1で対応づけ、Save をした選手がGKかどうかで判定する。
    GKは、試合JSONのポジション（GK）と、GK専用のイベント（Claim・Punch・KeeperPickup・KeeperSweeper）で判断する。
    events に blocked_by_outfield 列を付けて返す。対応する Save が見つからない SavedShot は、枠内として扱う。
    """
    d = _time_sorted(events)
    gk = {}
    if match_players is not None and len(match_players):
        mp = match_players[match_players["position_played"] == "GK"]
        for gid, pid in zip(mp["game_id"], mp["player_id"]):
            gk.setdefault(int(gid), set()).add(int(pid))
    only = d[d["type"].isin(GK_ONLY_EVENT_TYPES) & d["player_id"].notna()]
    for gid, pid in zip(only["game_id"], only["player_id"]):
        gk.setdefault(int(gid), set()).add(int(pid))

    blocked = pd.Series(False, index=d.index)
    n_gk = n_out = n_none = 0
    for gid, g in d.groupby("game_id", sort=False):
        typ = g["type"].to_numpy()
        team = g["team_id"].to_numpy()
        pid = g["player_id"].to_numpy()
        tt = g["_t"].to_numpy()
        idx = g.index.to_numpy()
        used = set()
        gks = gk.get(int(gid), set())
        for i in np.flatnonzero(typ == "SavedShot"):
            best = None
            for j in range(max(0, i - STOP_PAIR_WINDOW), min(len(g), i + STOP_PAIR_WINDOW + 1)):
                if (j != i and typ[j] == "Save" and team[j] != team[i] and j not in used
                        and abs(tt[j] - tt[i]) <= STOP_PAIR_MAX_GAP_S):
                    if best is None or abs(j - i) < abs(best - i) or (abs(j - i) == abs(best - i) and j > i):
                        best = j
            if best is None:
                n_none += 1
                continue
            used.add(best)
            if (not pd.isna(pid[best])) and int(pid[best]) in gks:
                n_gk += 1
            else:
                blocked.loc[idx[i]] = True
                n_out += 1
    out = _flat(events)
    out["blocked_by_outfield"] = blocked.reindex(out.index).fillna(False).astype(bool)
    n_saves = int((d["type"] == "Save").sum())
    print(f"  止められたシュート{n_gk + n_out + n_none}本の内訳: GKが止めた{n_gk} / 選手がブロック{n_out} / "
          f"対応するSaveなし{n_none}（Saveイベントは{n_saves}件）")
    return out


def compute_sca(events):
    """
    シュート・ゴールにつながったプレーを、選手ごと・試合ごとに数える（FBrefのSCA/GCAに近い近似）。
    各シュートについて、直前の同じチームのプレーを最大2つまでさかのぼり、次の種類でその選手に加算する。
      pass_live  … 流れの中の成功パス        pass_dead … セットプレーの成功パス（FK・CK・スローイン等）
      take_on    … ドリブル成功              shot      … 直前のシュート（こぼれ球からのシュート）
      foul_won   … ファウル獲得              defensive … タックル成功・インターセプト
    相手のプレーが挟まる、前半と後半をまたぐ、時間差が SCA_MAX_GAP_S 秒を超える、場合はそこで止める。
    PKのシュートとオウンゴールは、シュートに数えない。gca は、そのシュートがゴールだったもの。
    戻り値: game_id・player_id ごとの回数（sca_pass_live … sca, gca）。
    """
    d = _time_sorted(events)

    t = d["type"]
    ok = d["ok"]
    set_piece = d["q_ThrowIn"] | d["q_GoalKick"] | d["q_CornerTaken"] | d["q_FreekickTaken"]
    is_shot = d["is_shot"] & ~d["q_OwnGoal"] & ~d["q_Penalty"]
    kind = pd.Series(None, index=d.index, dtype=object)
    kind[(t == "Pass") & ok & ~set_piece] = "pass_live"
    kind[(t == "Pass") & ok & set_piece] = "pass_dead"
    kind[(t == "TakeOn") & ok] = "take_on"
    kind[(t == "Foul") & ok] = "foul_won"
    kind[((t == "Tackle") & ok) | (t == "Interception")] = "defensive"
    kind[is_shot] = "shot"
    d["_kind"] = kind

    rows = []
    for gid, g in d.groupby("game_id", sort=False):
        team = g["team_id"].to_numpy()
        typ = t.loc[g.index].to_numpy()
        okv = ok.loc[g.index].to_numpy()
        pid = g["player_id"].to_numpy()
        tt = g["_t"].to_numpy()
        per = g["_period"].to_numpy()
        kd = g["_kind"].to_numpy()
        shot = is_shot.loc[g.index].to_numpy()
        goal = (g["is_goal"] & ~g["q_OwnGoal"]).to_numpy()
        for i in np.flatnonzero(shot):
            credited, j = 0, i - 1
            while j >= 0 and credited < 2:
                if per[j] != per[i] or tt[i] - tt[j] > SCA_MAX_GAP_S:
                    break
                if team[j] != team[i]:
                    if typ[j] in SCA_BREAK_OPP or (typ[j] == "Tackle" and okv[j]):
                        break                                 # 相手がボールを持った（パス・奪取・タックル成功など）
                    j -= 1                                    # 相手の反則・競り合い・GKのセーブなどは、連なりを切らない
                    continue
                if isinstance(kd[j], str) and not pd.isna(pid[j]):
                    rows.append((gid, int(pid[j]), kd[j], bool(goal[i])))
                    credited += 1
                j -= 1
    cols = [f"sca_{k}" for k in SCA_KINDS] + ["sca", "gca"]
    if not rows:
        return pd.DataFrame(columns=["game_id", "player_id"] + cols)
    r = pd.DataFrame(rows, columns=["game_id", "player_id", "kind", "goal"])
    out = r.pivot_table(index=["game_id", "player_id"], columns="kind", values="goal", aggfunc="size", fill_value=0)
    out = out.reindex(columns=SCA_KINDS, fill_value=0)
    out.columns = [f"sca_{k}" for k in SCA_KINDS]
    out["sca"] = out.sum(axis=1)
    out["gca"] = r[r["goal"]].groupby(["game_id", "player_id"]).size().reindex(out.index).fillna(0).astype(int)
    return out.reset_index()


def compute_carries(spadl):
    """
    SPADL（socceraction）の「dribble」（連続する2つのプレーの間を補完したボール運び）から、
    選手ごと・試合ごとのボール運びを数える。座標は m、どちらのチームも左から右へ攻める向きにそろっている。
    progressive_carries … 前進10m以上、かつ自陣40%で終わらないもの（ペナルティエリアへの侵入も含める）
    """
    from socceraction.spadl import config as spadlconfig

    dr = spadl[spadl["type_id"] == spadlconfig.actiontypes.index("dribble")].copy()
    cols = ["carries", "carry_distance", "progressive_carries", "carries_final_third", "carries_into_box"]
    if dr.empty:
        return pd.DataFrame(columns=["game_id", "player_id"] + cols)
    dx = dr["end_x"] - dr["start_x"]
    dist = np.hypot(dx, dr["end_y"] - dr["start_y"])
    box_x, y_lo, y_hi = PITCH_L * BOX_X / 100, PITCH_W * BOX_Y_LO / 100, PITCH_W * BOX_Y_HI / 100
    end_box = (dr["end_x"] >= box_x) & dr["end_y"].between(y_lo, y_hi)
    start_box = (dr["start_x"] >= box_x) & dr["start_y"].between(y_lo, y_hi)
    dr["carries"] = 1
    dr["carry_distance"] = dist
    dr["progressive_carries"] = (((dx >= CARRY_PROGRESSIVE_MIN_M) & (dr["end_x"] >= PITCH_L * 0.4)) | (end_box & ~start_box))
    dr["carries_final_third"] = (dr["start_x"] < PITCH_L * FINAL_THIRD_X / 100) & (dr["end_x"] >= PITCH_L * FINAL_THIRD_X / 100)
    dr["carries_into_box"] = end_box & ~start_box
    g = dr.groupby(["game_id", "player_id"])[cols].sum().reset_index()
    g["carry_distance"] = g["carry_distance"].round(1)
    return g


def build_match_table(events, match_players, sca=None, carries=None):
    """
    選手×試合の表。試合JSONが読めれば、先発・出場時間・評価・試合最優秀選手も付く。
    sca / carries があれば、シュートにつながったプレー・ボール運びの回数も付ける。
    """
    keys = ["league", "season", "date", "game_id", "game", "team", "opponent", "venue",
            "player_id", "player"]
    keys += [c for c in ("goals_for", "goals_against") if c in events.columns]
    m = _int_ids(aggregate(events, by=keys).drop(columns="matches"))
    if match_players is not None and not match_players.empty:
        m = m.merge(_int_ids(match_players.copy()), on=["game_id", "player_id"], how="left")
        for c in ("is_starter", "man_of_match"):
            m[c] = m[c].map({True: 1, False: 0}).astype("Int64")
    for c in ("position_played", "is_starter", "man_of_match", "minutes", "rating"):
        if c not in m.columns:                           # 試合JSONが読めなかった場合も、列は空欄で作る
            m[c] = np.nan
    for extra, cols in ((sca, [f"sca_{k}" for k in SCA_KINDS] + ["sca", "gca"]),
                        (carries, ["carries", "carry_distance", "progressive_carries",
                                   "carries_final_third", "carries_into_box"])):
        if extra is None:
            continue
        if len(extra):
            m = m.merge(_int_ids(extra.copy()), on=["game_id", "player_id"], how="left")
            m[cols] = m[cols].fillna(0)
        else:                                            # 該当が1件も無い場合も、列は0で作っておく
            for c in cols:
                m[c] = 0
    m = _split_saves_and_gk(m)
    first = [c for c in ("position_played", "is_starter", "man_of_match", "minutes", "rating") if c in m.columns]
    rest = [c for c in m.columns if c not in first and c not in keys]
    m = m[keys + first + rest]
    return m.sort_values(["date", "game", "team", "player"]).reset_index(drop=True)


def _split_saves_and_gk(m):
    """
    GKの指標を作る。
    - 「Save」のイベントは、GKならセーブ（saves）、フィールドプレーヤーならシュートブロック（shot_blocks）に分ける。
    - GKには、失点（goals_conceded。自チームの選手のオウンゴールは除く）とクリーンシートを付ける。
      GKが複数いる試合は、出場時間の比で失点を配分する。クリーンシートは、無失点で60分以上出場したGK。
    """
    m = m.copy()
    pos = m["position_played"] if "position_played" in m else pd.Series("", index=m.index)
    gk_events = m[["claims", "punches", "keeper_pickups", "gk_sweeper_actions"]].sum(axis=1) > 0
    is_gk = (pos == "GK") | ((pos == "Sub") & gk_events)
    m["saves"] = np.where(is_gk, m["save_events"], 0)
    m["shot_blocks"] = np.where(is_gk, 0, m["save_events"])
    m = m.drop(columns="save_events")
    m["goals_conceded"] = np.nan
    m["clean_sheets"] = np.nan
    if "goals_against" in m and is_gk.any():
        own = m.groupby(["game_id", "team"])["own_goals"].transform("sum")
        conceded = (pd.to_numeric(m["goals_against"], errors="coerce") - own).clip(lower=0)
        w = pd.to_numeric(m["minutes"], errors="coerce") if "minutes" in m else pd.Series(np.nan, index=m.index)
        w = w.where(is_gk).fillna(0)
        n_gk = is_gk.groupby([m["game_id"], m["team"]]).transform("sum")
        tot_w = w.groupby([m["game_id"], m["team"]]).transform("sum")
        share = np.where(tot_w > 0, w / tot_w.replace(0, np.nan), 1 / n_gk.replace(0, np.nan))
        m.loc[is_gk, "goals_conceded"] = (conceded * share).round()[is_gk]
        clean = (pd.to_numeric(m["goals_against"], errors="coerce") == 0) & (w >= 60)
        m.loc[is_gk, "clean_sheets"] = clean[is_gk].astype(int)
    return m


def build_player_table(matches):
    """
    選手のシーズン集計を、試合別の表（matches）から作る。venue = all / home / away の行を縦に並べる。
    試合別の表から作るので、既存の試合に新しい試合を足したあとでも同じ手順で作り直せる。
    """
    base = ["league", "season", "team", "player_id", "player"]
    m = matches.copy()
    for c in ("is_starter", "man_of_match", "minutes", "rating"):
        if c not in m:
            m[c] = np.nan
        m[c] = pd.to_numeric(m[c], errors="coerce")
    sum_cols = [c for c in COUNT_COLS if c != "save_events"] + [c for c in EXTRA_COUNT_COLS if c in m.columns]
    gk_cols = [c for c in ("goals_conceded", "clean_sheets") if c in m.columns]   # GK以外は空欄のままにする
    frames = []
    for venue in ("all", "home", "away"):
        mt = m if venue == "all" else m[m["venue"] == venue]
        if mt.empty:
            continue
        g = mt.groupby(base, dropna=False)
        out = g[[c for c in sum_cols if c not in gk_cols]].sum()
        for c in out.columns:
            if c != "carry_distance":
                out[c] = out[c].astype(int)
        for c in gk_cols:
            out[c] = g[c].sum(min_count=1)
        out.insert(0, "matches", g.size())
        out.insert(1, "starts", g["is_starter"].sum(min_count=1))
        out.insert(2, "minutes", g["minutes"].sum(min_count=1))
        out.insert(3, "motm", g["man_of_match"].sum(min_count=1))
        out.insert(4, "rating_avg", g["rating"].mean().round(2))
        out = _add_rates(out).reset_index()
        out.insert(2, "venue", venue)
        frames.append(out)
    return _int_ids(pd.concat(frames, ignore_index=True))


# ---- Understat（xG・xA）との結合 --------------------------------------------------
def norm_name(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return " ".join(s.replace("-", " ").split())


US_COLS = {  # Understatの列 -> 出力名
    "matches": "us_apps", "minutes": "us_minutes",
    "goals": "us_goals", "assists": "us_assists",
    "xg": "us_xg", "np_xg": "us_np_xg", "xa": "us_xa",
    "xg_chain": "us_xg_chain", "xg_buildup": "us_xg_buildup",
    "yellow_cards": "us_yellow_cards", "red_cards": "us_red_cards",
}


def add_understat(stats, us):
    """
    選手のシーズン表に、Understatの出場数・出場時間・xG等を選手名で結合する（venue=all の行のみ）。
    出場時間（minutes）が試合JSONから取れていなければ、Understatの値で埋める。
    未マッチの選手（名前の表記ゆれなど）は空欄になる。
    """
    u = _flat(us)
    u["_key"] = u["player"].map(norm_name)
    u = (u.groupby("_key", as_index=False)[list(US_COLS)].sum().rename(columns=US_COLS))
    s = stats.assign(_key=stats["player"].map(norm_name)).merge(u, on="_key", how="left")
    s = s.drop(columns="_key")
    us_cols = list(US_COLS.values())
    if "venue" in s:
        s.loc[s["venue"] != "all", us_cols] = np.nan
    s["xg_diff"] = (s["goals"] - s["us_xg"]).round(2)          # ゴール − xG（venue=all の行のみ）
    if "minutes" not in s or s["minutes"].isna().all():
        s["minutes"] = s["us_minutes"]
    miss = s.loc[(s.get("venue", "all") == "all") & s["us_minutes"].isna(), "player"].unique()
    if len(miss):
        print(f"Understatと未マッチ: {len(miss)}人（例: {list(miss[:5])}）")
    return s


_P90_BASE = ["passes", "progressive_passes", "passes_final_third", "passes_into_box", "key_passes",
             "crosses", "long_balls", "through_balls", "take_ons_won", "dispossessed", "errors",
             "tackles", "tackles_won", "interceptions", "clearances", "blocked_passes", "ball_recoveries",
             "def_actions_att_third", "aerials_won", "fouls_committed", "fouls_won", "dribbled_past",
             "shots", "shots_on_target", "goals", "touches", "touches_att_pen",
             "carries", "progressive_carries", "sca", "gca", "saves", "goals_conceded",
             "us_xg", "us_xa", "xg_diff"]
DEFAULT_P90 = list(dict.fromkeys(
    _P90_BASE + [c[:-4] for specs in RANK_SPECS.values() for c, *_ in specs if c.endswith("_p90")]))


def add_per90(stats, cols=None):
    """minutes 列を使って90分あたりに換算する（存在する列のみ。出場時間が短い行は空欄）。"""
    s = stats.copy()
    if "minutes" not in s:
        return s
    mins = pd.to_numeric(s["minutes"], errors="coerce")
    mins = mins.where(mins >= MIN_MINUTES_P90)
    new = {c + "_p90": (s[c] / mins * 90).round(2) for c in (cols or DEFAULT_P90) if c in s.columns}
    return pd.concat([s, pd.DataFrame(new, index=s.index)], axis=1)


# ---- ポジションと順位 ----------------------------------------------------------
def add_positions(players, matches):
    """
    選手ごとに position（FW/MF/DF/GK）と position_detail（DFはCB/SB、それ以外は position と同じ）を付ける。
    試合ごとのポジション（position_played）のうち、出場時間が最も長い系統を採用する。
    交代で入った試合（Sub）は判定に使わない。Subでしか出ていない選手は空欄になる。
    """
    m = matches[["team", "player_id", "position_played", "minutes"]].copy()
    m["grp"] = m["position_played"].map(POS_GROUP)
    m = m[m["grp"].notna()]
    m["w"] = pd.to_numeric(m["minutes"], errors="coerce")
    if m["w"].isna().all():
        m["w"] = 1.0                                        # 出場時間が不明なら試合数で数える
    m["w"] = m["w"].fillna(0)
    keys = ["team", "player_id"]
    p = players.copy()
    if m.empty:
        p["position"], p["position_detail"] = np.nan, np.nan
        return p
    g = m.groupby(keys + ["grp"])["w"].sum().reset_index()
    main = g.loc[g.groupby(keys)["w"].idxmax(), keys + ["grp"]].rename(columns={"grp": "position"})
    df_ = m[m["grp"] == "DF"]
    tot = df_.groupby(keys)["w"].sum()
    side = df_[df_["position_played"].isin(SIDE_BACK_POSITIONS)].groupby(keys)["w"].sum()
    share = (side.reindex(tot.index).fillna(0) / tot.replace(0, np.nan)).rename("_sb_share").reset_index()
    main = main.merge(share, on=keys, how="left")
    main["position_detail"] = np.where(
        main["position"] == "DF", np.where(main["_sb_share"].fillna(0) >= SB_MIN_SHARE, "SB", "CB"),
        main["position"])
    p = p.merge(main[keys + ["position", "position_detail"]], on=keys, how="left")
    cols = [c for c in p.columns if c not in ("position", "position_detail")]
    i = cols.index("player") + 1
    return p[cols[:i] + ["position", "position_detail"] + cols[i:]]


def add_ranks(players, min_minutes=RANK_MIN_MINUTES):
    """
    RANK_SPECS の指標について、順位（<指標>_順位）と母数（<指標>_順位_母数）を付ける。
    順位は「同じリーグ・同じシーズン・同じ集団（FW / MF / CB / SB / GK）・同じ venue」の中で付け、
    シーズン通算の出場時間が min_minutes 以上で、分母（試行数）が下限以上の選手だけが対象になる。
    条件を満たさない行は、順位・母数とも空欄。母数は、順位を付けた対象の人数。1位が最も良い。
    venue が home / away の行は、その venue の値で順位を付ける（出場時間の条件は通算で判定する）。
    """
    p = players.copy()
    if "position" not in p or p["position"].isna().all():
        return p
    keyc = ["league", "season", "team", "player_id"]
    base = p[p["venue"] == "all"].set_index(keyc)
    idx = pd.MultiIndex.from_frame(p[keyc])

    def from_all(series):
        return pd.Series(series.reindex(idx).to_numpy(), index=p.index)

    tot_min = from_all(pd.to_numeric(base["minutes"], errors="coerce"))
    derived = {                                   # 列としては無い分母（通算の値から作る）
        "challenges": lambda b: b["tackles"] + b["dribbled_past"],
        "duels": lambda b: b["tackles"] + b["aerials"],
        "shots_faced": lambda b: b["saves"] + b["goals_conceded"],
    }
    denoms = {}
    for specs in RANK_SPECS.values():
        for _, _, denom, _ in specs:
            if denom and denom not in denoms:
                try:
                    denoms[denom] = from_all(derived[denom](base) if denom in derived else base[denom])
                except KeyError:
                    denoms[denom] = None          # その分母の列が無い（例: ボール運びを取っていない）→ 指標ごと飛ばす
    group = pd.Series(np.where(p["position"] == "DF", p["position_detail"], p["position"]), index=p.index)

    new = {}
    for key, specs in RANK_SPECS.items():
        for col, direction, denom, dmin in specs:
            if col not in p:
                continue
            ok = (group == key) & (tot_min >= min_minutes) & p[col].notna()
            if denom:
                if denoms.get(denom) is None:
                    continue
                ok &= denoms[denom].fillna(0) >= dmin
            rk, pop = f"{col}_順位", f"{col}_順位_母数"
            if rk not in new:
                new[rk] = pd.Series(pd.array([pd.NA] * len(p), dtype="Int64"), index=p.index)
                new[pop] = pd.Series(pd.array([pd.NA] * len(p), dtype="Int64"), index=p.index)
            if not ok.any():
                continue
            grp = p.loc[ok].groupby(["league", "season", "venue"])[col]
            new[rk].loc[ok] = grp.rank(ascending=(direction == "low"), method="min").astype(int)
            new[pop].loc[ok] = grp.transform("size").astype(int)
    return pd.concat([p, pd.DataFrame(new, index=p.index)], axis=1)


# ---- Excel出力 --------------------------------------------------------------
def _format_sheet(ws, df):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def write_workbook(path, sheets, note=None):
    """{シート名: DataFrame} を、渡した順にxlsxへ書く。note はファイルのプロパティ（説明）に記録する。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _format_sheet(xw.sheets[name[:31]], df)
        if note:
            xw.book.properties.description = note


EVENT_SHEETS = {"players", "matches"}
COMBINED_DIR = "結合"
SEASON_FILE_RE = re.compile(r"^.+_(\d{4})\.xlsx$")       # LaLiga_2025.xlsx など
SD_TO_SHORT = {v: k for k, v in SD_LEAGUES.items()}       # 旧ファイルのリーグ名（ENG-Premier League 等）の読み替え


def save_job(out, name, year, players, matches):
    """<out>/<リーグ名>/<リーグ名>_<開幕年>.xlsx に保存する。別スクリプトのxlsxがあれば上書きせずに止める。"""
    path = Path(out) / name / f"{name}_{year}.xlsx"
    if path.exists():
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
        if not sheets <= EVENT_SHEETS:
            raise RuntimeError(f"{path} は別のスクリプト（表の取得）のファイルのようです。"
                               "上書きを避けるため、--out を別のフォルダにしてください")
    write_workbook(path, {"players": players, "matches": matches}, note=VERSION_NOTE)
    return path


_QUALIFIERS_CHECKED = False


def output_path(out, name, year):
    return Path(out) / name / f"{name}_{year}.xlsx"


def load_existing_matches(path):
    """
    既存のxlsxから matches シートを読む（続きから追加するため）。
    ファイルが無い／古い形式（版が違う）の場合は None を返し、作り直しになる。
    """
    path = Path(path)
    if not path.exists():
        return None
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    sheets, note = set(wb.sheetnames), wb.properties.description
    wb.close()
    if not sheets <= EVENT_SHEETS:
        raise RuntimeError(f"{path} は別のスクリプト（表の取得）のファイルのようです。"
                           "上書きを避けるため、--out を別のフォルダにしてください")
    if note != VERSION_NOTE or "matches" not in sheets:
        print("  既存のファイルは古い形式のため、作り直します")
        return None
    return _int_ids(pd.read_excel(path, sheet_name="matches", dtype={"player_id": str}))


def process_job(name, year, out, n=None, last=False, rebuild=False, use_understat=True,
                use_approx=True, fetch=None, fetch_us=None):
    """
    1リーグ・1シーズンを取得して保存する。既存のxlsxがあれば、その matches に新しい試合の行を足す
    （取得済みの試合は取り直さない）。players は、足したあとの matches から作り直す。
    戻り値: (players, matches, 保存先, 新しく足した試合数)
    """
    fetch = fetch or fetch_events
    fetch_us = fetch_us or fetch_understat
    path = output_path(out, name, year)
    existing = None if rebuild else load_existing_matches(path)
    skip = set(existing["game_id"].dropna().astype(int)) if existing is not None else set()

    kwargs = {} if use_approx else {"with_carries": False}      # 近似の指標（ボール運び）を取らない
    events, schedule, files, carries = fetch(SD_LEAGUES[name], year, n, last, skip, **kwargs)
    new, n_new = None, 0
    if events is not None:
        global _QUALIFIERS_CHECKED
        if not _QUALIFIERS_CHECKED:
            check_qualifiers(events)                   # 初回はqualifier名を確認
            _QUALIFIERS_CHECKED = True
        events = add_match_info(events, schedule)
        mp = read_match_players(files)
        events = mark_shot_stops(events, mp)              # SavedShot を、GKのセーブとブロックに分ける
        if mp.empty or mp["rating"].isna().all():
            print("  注意: 試合JSONから評価・出場時間を読めませんでした（rating等は空欄になります）")
        sca = compute_sca(events) if use_approx else None   # シュートにつながったプレー（近似）
        new = build_match_table(events, mp, sca=sca, carries=carries)
        new["league"], new["season"] = name, season_label(year)
        n_new = new["game_id"].nunique()

    parts = [d for d in (existing, new) if d is not None and len(d)]
    if not parts:
        raise RuntimeError("試合データがありません")
    matches = (pd.concat(parts, ignore_index=True)
               .drop_duplicates(["game_id", "player_id"], keep="last")
               .sort_values(["date", "game", "team", "player"]).reset_index(drop=True))
    players = add_positions(build_player_table(matches), matches)
    if use_understat and n is None:
        try:
            players = add_understat(players, fetch_us(SD_LEAGUES[name], year))
        except Exception as e:  # noqa: BLE001
            print(f"  Understatの結合をスキップ: {e}")
    elif n is not None:
        print("  注意: --n 指定のため Understat は結合しません（動作確認用）")
    players = add_ranks(add_per90(players))          # 順位は同じリーグの中で付ける
    players["league"], players["season"] = name, season_label(year)
    return players, matches, save_job(out, name, year, players, matches), n_new


def _event_files(out):
    """<out>/<リーグ>/<リーグ>_<年>.xlsx のうち、イベント集計のファイル（players・matchesシート）だけを返す。"""
    import openpyxl

    out = Path(out)
    found = []
    for f in out.glob("*/*.xlsx"):
        m = SEASON_FILE_RE.match(f.name)
        if f.parent.name == COMBINED_DIR or f.name.startswith("~$") or not m:
            continue
        wb = openpyxl.load_workbook(f, read_only=True)
        ok = set(wb.sheetnames) == EVENT_SHEETS
        wb.close()
        if ok:
            found.append((f, int(m[1])))
    order = {k: i for i, k in enumerate(SD_LEAGUES)}
    return sorted(found, key=lambda t: (t[1], order.get(t[0].parent.name, 99)))


def write_by_season(d, kind, sheets):
    """
    結合ファイルをシーズンごとに分けたxlsxも作る（結合/<開幕年>年/all_leagues_<players|matches>_<開幕年>.xlsx）。
    matches は全リーグ・1シーズンで約19MB（GitHubのブラウザアップロードは1ファイル25MiBまで）。
    書式は付けない。xlsxwriter があればそれで書く（無ければ openpyxl。約2割大きくなる）。player_id は数値に戻す。
    """
    try:
        import xlsxwriter  # noqa: F401
        engine = "xlsxwriter"                       # openpyxl より約2割小さいファイルになる（matches 約19MB → 約23MB）
    except ImportError:
        engine = "openpyxl"
        print("  ※ xlsxwriter が無いため、結合ファイルが大きめになります（python -m pip install xlsxwriter で入ります）")
    for label, df in sheets.items():
        year = label[:4]
        path = Path(d) / f"{year}年" / f"all_leagues_{kind}_{year}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        df = df.assign(player_id=pd.to_numeric(df["player_id"], errors="coerce").astype("Int64"))
        with pd.ExcelWriter(path, engine=engine) as xw:
            df.to_excel(xw, sheet_name=label[:31], index=False)
        print(f"  結合（{kind}）: {label} → {path}（{path.stat().st_size / 1048576:.1f}MiB）")


def rebuild_combined(out):
    """
    全リーグのファイル（過去の実行分も含む）を、シーズンごとに結合して、次のxlsxを作り直す（GitHubへブラウザで上げる用に、シーズン別に分けている）。
      結合/<開幕年>年/all_leagues_players_<開幕年>.xlsx  … 選手のシーズン集計（全リーグ・1シーズン。1シート）
      結合/<開幕年>年/all_leagues_matches_<開幕年>.xlsx  … 選手×試合の表（同上）
    """
    files = _event_files(out)
    if not files:
        print("結合するファイルがありません")
        return None
    by_kind = {"players": {}, "matches": {}}
    for f, year in files:
        label = season_label(year)
        book = pd.read_excel(f, sheet_name=list(EVENT_SHEETS), dtype={"player_id": str})
        for kind, df in book.items():
            df["league"] = df["league"].replace(SD_TO_SHORT)
            df["season"] = label
            by_kind[kind].setdefault(label, []).append(df)
    d = Path(out) / COMBINED_DIR
    for kind, seasons in by_kind.items():
        sheets = {k: pd.concat(seasons[k], ignore_index=True) for k in sorted(seasons, reverse=True)}
        rows = {k: len(v) for k, v in sheets.items()}
        print(f"結合（{kind}）: {len(files)}ファイル  シーズン別の行数: {rows}")
        write_by_season(d, kind, sheets)
    return d


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--leagues", nargs="*", help="リーグ名（部分一致・複数可。省略で全リーグ）")
    p.add_argument("--seasons", nargs="*", help="シーズン（開幕年。省略で最新）")
    p.add_argument("--n", type=int, default=None, help="取得する試合数（省略で終了した全試合）")
    p.add_argument("--last", action="store_true", help="開幕からではなく直近のn試合を取る")
    p.add_argument("--rebuild", action="store_true",
                   help="既存のファイルを使わず、最初から作り直す（取得済みの試合は保存分を使うので速い）")
    p.add_argument("--no-understat", action="store_true", help="Understat（xG・xA）を結合しない")
    p.add_argument("--no-approx", action="store_true",
                   help="近似を含む指標（シュートにつながったプレー・ボール運び）を作らない")
    g = p.add_mutually_exclusive_group()
    g.add_argument("--no-combine", action="store_true",
                   help="取得のあとに結合しない（並列で実行するときに使い、最後に --combine-only で1回だけ結合する）")
    g.add_argument("--combine-only", action="store_true",
                   help="取得はせず、出力フォルダにあるファイルの結合だけを行う")
    p.add_argument("--out", default=str(DEFAULT_OUT),
                   help="出力フォルダ（省略時はこのスクリプトと同じ場所の output フォルダ）")
    args = p.parse_args()

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise SystemExit("Excel出力に openpyxl が必要です: python -m pip install openpyxl")

    print(f"出力先: {Path(args.out).resolve()}")
    if args.combine_only:
        rebuild_combined(args.out)
        return
    saved = 0
    for name, year in build_jobs(args.leagues, args.seasons):
        label = f"{name} {year}"
        try:
            players, matches, path, n_new = process_job(
                name, year, args.out, n=args.n, last=args.last, rebuild=args.rebuild,
                use_understat=not args.no_understat, use_approx=not args.no_approx)
        except PermissionError:
            print(f"[{label}] xlsxに保存できません（Excelで開いていれば閉じて再実行してください）")
            continue
        except Exception as e:  # noqa: BLE001
            print(f"[{label}] 失敗: {e}")
            continue
        print(f"[{label}] 新しく{n_new}試合を追加 → 合計{matches['game_id'].nunique()}試合・"
              f"選手{players['player_id'].nunique()}人・{len(matches)}行（選手×試合）→ {path}")
        saved += 1
    if saved and args.no_combine:
        print("結合はしませんでした（--no-combine）。あとで --combine-only で結合してください")
    elif saved:
        try:
            rebuild_combined(args.out)
        except PermissionError:
            print("結合ファイルを保存できません（Excelで開いている場合は閉じて、--combine-only で再実行してください）")


if __name__ == "__main__":
    main()